#!/usr/bin/env python3
"""For every risk in the register, list the Annex A controls that treat it and
their implementation status in the Statement of Applicability.

Usage: python3 control_status.py ../risk-register/saarland-university-risk-register.xlsx \
           ../soa/saarland-university-soa.xlsx [out.md]
Needs openpyxl. Statuses in the SoA are invented portfolio data.
"""
import re
import sys

import openpyxl


def main(reg_path, soa_path, out="control-status.md"):
    soa = {}
    ws = openpyxl.load_workbook(soa_path)["SoA"]
    for r in ws.iter_rows(min_row=5, values_only=True):
        if r[0]:
            soa[str(r[0]).strip()] = (r[3], r[9])
    ws = openpyxl.load_workbook(reg_path)["Risk Register"]
    rows = []
    for r in ws.iter_rows(min_row=5, values_only=True):
        if not r[0]:
            continue
        ctrls = re.findall(r"(A\.\d+\.\d+)", r[11] or "")
        st = [soa.get(c, ("?", "?"))[1] for c in ctrls]
        rows.append((r[0], r[1], r[6], r[7], ctrls, st))
    weak_states = {"Planned", "Not implemented"}
    L = ["# Controls treating each risk and their SoA status", "",
         "Generated from the risk register and the Statement of Applicability. Statuses are invented portfolio data. "
         "'Weak' counts controls that are Planned or Not implemented.", "",
         "| Risk | L × I | Controls (SoA status) | Weak controls |", "|---|---|---|---|"]
    for rid, title, l, i, ctrls, st in rows:
        weak = sum(s in weak_states for s in st)
        cs = "; ".join(f"{c} ({s})" for c, s in zip(ctrls, st))
        L.append(f"| {rid} {title} | {l}×{i} | {cs} | {weak} of {len(ctrls)} |")
    open(out, "w").write("\n".join(L) + "\n")
    print("\n".join(L))


if __name__ == "__main__":
    main(*sys.argv[1:4])
